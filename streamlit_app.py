import streamlit as st
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from io import BytesIO
from PIL import Image
import pytesseract
from deep_translator import GoogleTranslator

# ============================
# EMBEDDED LOGO (BASE64 STRING)
# ============================
RAJYUG_LOGO_BASE64 = """
PUT_YOUR_LOGO_BASE64_HERE
"""

# ----------------------------
# Extract text + formatting from PDF
# ----------------------------
def extract_text_from_pdf(pdf_document):
    elements = []
    for page_num in range(len(pdf_document)):
        page = pdf_document.load_page(page_num)
        page_width = page.rect.width

        text_blocks = page.get_text("dict")["blocks"]
        for block in text_blocks:
            if block["type"] == 0:
                for line in block["lines"]:
                    line_text = "".join([span["text"] for span in line["spans"]])

                    font_size = line["spans"][0]["size"] if line["spans"] else None
                    font_name = line["spans"][0]["font"] if line["spans"] else None
                    is_bold = "Bold" in font_name if font_name else False

                    x0 = line["bbox"][0]
                    if x0 < page_width * 0.3:
                        alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
                    elif x0 > page_width * 0.7:
                        alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
                    else:
                        alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

                    elements.append(('text', line_text, font_name, font_size, is_bold, alignment))

    return elements

# ----------------------------
# OCR from image
# ----------------------------
def extract_text_from_image(image):
    return pytesseract.image_to_string(image)

# ----------------------------
# Create DOCX output
# ----------------------------
def create_docx_with_formatting(elements):
    doc = Document()

    for element_type, *data in elements:
        if element_type == "text":
            text, font_name, font_size, is_bold, alignment = data

            # FIX: Make text always safe
            safe_text = str(text or "").strip()

            paragraph = doc.add_paragraph()
            run = paragraph.add_run(safe_text)

            if font_size:
                run.font.size = Pt(font_size)

            if font_name:
                run.font.name = font_name

            run.bold = is_bold
            paragraph.alignment = alignment

    doc_io = BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

# ----------------------------
# Translate using deep-translator
# ----------------------------
def translate_text(text, target_language):
    try:
        return GoogleTranslator(source='auto', target=target_language).translate(text)
    except Exception as e:
        return f"Translation error: {e}"

# =====================================
# STREAMLIT APP
# =====================================
def main():

    # Logo Display
    st.markdown(
        f"""
        <div style='text-align: center; padding-bottom: 30px;'>
            <img src="data:image/jpeg;base64,{RAJYUG_LOGO_BASE64}" width="250">
        </div>
        """,
        unsafe_allow_html=True
    )

    st.sidebar.title("PDF & Image Translator")
    st.sidebar.write("Upload → Extract → Translate → Download DOCX")

    uploaded_file = st.sidebar.file_uploader("Upload File", type=["pdf", "jpg", "jpeg", "png"])

    if uploaded_file:
        if uploaded_file.type == "application/pdf":
            st.info("Extracting text from PDF...")
            pdf_document = fitz.open(stream=uploaded_file.read(), filetype="pdf")
            all_elements = extract_text_from_pdf(pdf_document)

        else:
            st.info("Extracting text from Image...")
            image = Image.open(uploaded_file)
            text = extract_text_from_image(image)
            all_elements = [('text', text, None, None, False, WD_PARAGRAPH_ALIGNMENT.LEFT)]

        target_language = st.sidebar.selectbox("Select Target Language", ["Marathi", "Hindi"])
        language_code = "mr" if target_language == "Marathi" else "hi"

        translated_elements = []
        for element_type, *data in all_elements:
            if element_type == "text":
                text = data[0]
                translated_text = translate_text(text, language_code)
                translated_elements.append(('text', translated_text, *data[1:]))

        st.success(f"Content Translated to {target_language}")

        st.info("Generating DOCX...")
        docx_file = create_docx_with_formatting(translated_elements)

        st.download_button(
            label=f"Download DOCX ({target_language})",
            data=docx_file,
            file_name=f"translated_{target_language.lower()}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    else:
        st.warning("Please upload a PDF or image file.")

    # Footer
    st.markdown(
        """
        <style>
        .footer {
            position: fixed;
            bottom: 0;
            left: 170px;
            width: 100%;
            background: #000;
            color: white;
            text-align: center;
            padding: 10px;
        }
        </style>
        <div class="footer">
            Developed by @Rajyug IT Solutions Pvt. Ltd
        </div>
        """,
        unsafe_allow_html=True,
    )

if __name__ == "__main__":
    main()
